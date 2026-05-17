import hashlib
import os
import re
from difflib import SequenceMatcher

import PyPDF2
import docx
import openpyxl
import xlrd


def generate_sha256_from_path(file_path):
    sha256 = hashlib.sha256()

    with open(file_path, "rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            sha256.update(chunk)

    return sha256.hexdigest()


def normalize_text(text):
    text = text.lower()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"[^\w\s]", "", text)
    return text.strip()


def similarity_percent(text1, text2):
    text1 = normalize_text(text1)
    text2 = normalize_text(text2)

    if not text1 or not text2:
        return 0

    return SequenceMatcher(None, text1, text2).ratio() * 100


def extract_text_from_path(file_path):
    extension = os.path.splitext(file_path)[1].lower()

    if extension == ".pdf":
        return extract_pdf_text(file_path)

    if extension == ".docx":
        return extract_docx_text(file_path)

    if extension == ".xlsx":
        return extract_xlsx_text(file_path)

    if extension == ".xls":
        return extract_xls_text(file_path)

    return extract_plain_text(file_path)


def extract_plain_text(file_path):
    with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
        return file.read()


def extract_pdf_text(file_path):
    text = ""

    with open(file_path, "rb") as file:
        reader = PyPDF2.PdfReader(file)

        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"

    return text


def extract_docx_text(file_path):
    text = ""

    document = docx.Document(file_path)

    for paragraph in document.paragraphs:
        text += paragraph.text + "\n"

    return text


def extract_xlsx_text(file_path):
    text = ""

    workbook = openpyxl.load_workbook(file_path, data_only=True, read_only=True)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            text += " ".join(str(cell) for cell in row if cell is not None) + "\n"

    workbook.close()
    return text


def extract_xls_text(file_path):
    text = ""

    workbook = xlrd.open_workbook(file_path)

    for sheet in workbook.sheets():
        for row_index in range(sheet.nrows):
            values = sheet.row_values(row_index)
            text += " ".join(str(value) for value in values if value) + "\n"

    return text